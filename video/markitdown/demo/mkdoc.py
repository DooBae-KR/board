import sys; sys.path.insert(0,'pylib')
from docx import Document
def make(path, title, bullets, head, rows):
    d=Document(); d.add_heading(title,1)
    for b in bullets: d.add_paragraph(b, style='List Bullet')
    t=d.add_table(rows=1, cols=len(head)); t.style='Table Grid'
    for i,h in enumerate(head): t.rows[0].cells[i].text=h
    for r in rows:
        c=t.add_row().cells
        for i,v in enumerate(r): c[i].text=v
    d.save(path)
make('report_ko.docx','3분기 보고서',['매출은 전 분기보다 12% 늘었어요','신규 고객 340명 확보'],['항목','금액','증감'],[['매출','1,240만원','+12%'],['비용','860만원','+3%']])
make('report_en.docx','Q3 Report',['Revenue grew 12% quarter over quarter','340 new customers acquired'],['Item','Amount','Change'],[['Revenue','$12.4k','+12%'],['Cost','$8.6k','+3%']])
